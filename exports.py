import csv
import os
import tempfile
from contextlib import contextmanager
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO, Iterator

from offer_page import CONTACT_FIELDS, SETTING_KEY as OFFER_PAGE_SETTING, parse_offer_page
from quote_domain import DEFAULT_OPERATION, OPERATIONS

def pl(value: Any) -> str:
    return f"{Decimal(value):.2f}".replace(".", ",")


PHOTO_WIDTH_MIN_MM = 25
PHOTO_WIDTH_MAX_MM = 56
PHOTO_WIDTH_DEFAULT_MM = 44


def photo_width_mm(settings: dict[str, Any]) -> float:
    try:
        value = float(str(settings.get("photo_width_mm", "")).replace(",", "."))
    except ValueError:
        return PHOTO_WIDTH_DEFAULT_MM
    return min(max(value, PHOTO_WIDTH_MIN_MM), PHOTO_WIDTH_MAX_MM)


WZ_ROW_HEIGHT_MIN_MM = 5
WZ_ROW_HEIGHT_MAX_MM = 9
WZ_ROW_HEIGHT_DEFAULT_MM = 5


def wz_row_height_mm(settings: dict[str, Any]) -> float:
    try:
        value = float(str(settings.get("wz_row_height_mm", "")).replace(",", "."))
    except ValueError:
        return WZ_ROW_HEIGHT_DEFAULT_MM
    return min(max(value, WZ_ROW_HEIGHT_MIN_MM), WZ_ROW_HEIGHT_MAX_MM)


@contextmanager
def _atomic_destination(destination: str | Path) -> Iterator[Path]:
    target = Path(destination)
    with tempfile.NamedTemporaryFile(
        prefix=f".{target.name}.",
        suffix=".tmp",
        dir=target.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)

    try:
        yield temporary_path
        temporary_path.replace(target)
    finally:
        temporary_path.unlink(missing_ok=True)


def _safe_spreadsheet_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    if value.startswith(("\t", "\r")) or value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _estimate_export_table(
    estimate: dict[str, Any],
    settings: dict[str, str] | None = None,
) -> tuple[list[str], list[list[Any]]]:
    header = estimate["header"]
    company = settings or {}
    fields = [
        "Numer wyceny",
        "Data",
        "Sprzedawca",
        "NIP sprzedawcy",
        "Adres sprzedawcy",
        "Telefon sprzedawcy",
        "E-mail sprzedawcy",
        "Strona sprzedawcy",
        "Zleceniodawca",
        "NIP klienta",
        "Adres klienta",
        "Numer zamówienia",
        "Uwagi",
        "Pozycja",
        "Operacja",
        "Zdjęcie poglądowe",
        "Materiał",
        "Grubość (mm)",
        "Sposób kalkulacji",
        "Ilość / zakres",
        "Cena materiału (zł/kg)",
        "Masa (kg/szt. lub razem)",
        "Czas palenia (sek./szt. lub razem)",
        "Stawka palenia (zł/godz.)",
        "VAT (%)",
        "Materiał netto / szt. (zł)",
        "Palenie netto / szt. (zł)",
        "Cena netto / szt. (zł)",
        "VAT / szt. (zł)",
        "Cena brutto / szt. (zł)",
        "Materiał netto razem (zł)",
        "Palenie netto razem (zł)",
        "Wartość netto (zł)",
        "VAT razem (zł)",
        "Wartość brutto (zł)",
    ]
    created_at = datetime.fromisoformat(header["created_at"]).astimezone().strftime("%Y-%m-%d %H:%M")
    rows: list[list[Any]] = []
    for item in estimate["items"]:
        whole_order = item.get("calculation_mode", "per_piece") == "whole_order"
        quantity = int(item["quantity"])
        material_total = Decimal(str(item["material_cost"])) * (1 if whole_order else quantity)
        burning_total = Decimal(str(item["burning_cost"])) * (1 if whole_order else quantity)
        unit_values = ("", "", "", "", "") if whole_order else (
            item["material_cost"],
            item["burning_cost"],
            item["unit_net"],
            item["unit_vat"],
            item["unit_gross"],
        )
        rows.append([
            header["quote_number"],
            created_at,
            company.get("company_name", ""),
            company.get("company_tax_id", ""),
            company.get("company_address", ""),
            company.get("company_phone", ""),
            company.get("company_email", ""),
            company.get("company_website", ""),
            header["customer_name"],
            header.get("customer_tax_id", ""),
            header.get("customer_address", ""),
            header["order_number"],
            header.get("notes", ""),
            item["element_name"],
            OPERATIONS.get(item.get("operation", DEFAULT_OPERATION), ""),
            "Tak" if item.get("preview_image") else "Nie",
            item["material"],
            item["thickness_mm"],
            "Cała partia" if whole_order else "Na sztukę",
            quantity,
            item["price_per_kg"],
            item["used_weight_kg"],
            item["duration_seconds"],
            item["burning_rate_per_hour"],
            item["vat_percent"],
            *unit_values,
            material_total,
            burning_total,
            item["line_net"],
            item["line_vat"],
            item["line_gross"],
        ])
    return fields, rows


def export_estimate_csv(
    estimate: dict[str, Any],
    destination: str | Path,
    settings: dict[str, str] | None = None,
) -> None:
    fields, rows = _estimate_export_table(estimate, settings)
    with _atomic_destination(destination) as temporary_path:
        with temporary_path.open("w", encoding="utf-8-sig", newline="") as csv_file:
            writer = csv.writer(csv_file, delimiter=";")
            writer.writerow(fields)
            for row in rows:
                writer.writerow(_safe_spreadsheet_value(value) for value in row)


def _excel_cell_value(value: Any, numeric: bool) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if numeric and isinstance(value, str) and value.strip():
        try:
            return float(value.replace(",", "."))
        except ValueError:
            pass
    return _safe_spreadsheet_value(value)


def _write_xlsx(
    destination: str | Path,
    sheet_title: str,
    header_row: list[str],
    rows: list[list[Any]],
    numeric_columns: set[int],
    widths: list[int] | None = None,
    leading_rows: list[list[Any]] | None = None,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise ImportError(
            "Eksport Excel wymaga biblioteki openpyxl. Zainstaluj zależności poleceniem: "
            "python -m pip install -r requirements.txt"
        ) from error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    for lead in leading_rows or []:
        sheet.append([_excel_cell_value(value, False) for value in lead])
    sheet.append(header_row)
    header_index = sheet.max_row
    for cell in sheet[header_index]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243F82")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in rows:
        sheet.append([
            _excel_cell_value(value, index in numeric_columns)
            for index, value in enumerate(row)
        ])
    for index in range(len(header_row)):
        width = widths[index] if widths and index < len(widths) else 16
        sheet.column_dimensions[get_column_letter(index + 1)].width = width
    sheet.freeze_panes = sheet.cell(row=header_index + 1, column=1)
    with _atomic_destination(destination) as temporary_path:
        workbook.save(temporary_path)


def export_estimate_xlsx(
    estimate: dict[str, Any],
    destination: str | Path,
    settings: dict[str, str] | None = None,
) -> None:
    fields, rows = _estimate_export_table(estimate, settings)
    numeric_columns = {17, 19, 20, 21, 22, 23, 24, *range(25, 35)}
    _write_xlsx(
        destination,
        "Wycena",
        fields,
        rows,
        numeric_columns,
        widths=[16, 17, 24, 14, 28, 16, 24, 20, 26, 14, 28, 18, 28, 26, 18, 12, 16, 12, 16, 12, 14, 14, 16, 14, 10, *([16] * 10)],
    )


def export_delivery_note_xlsx(
    delivery_note: dict[str, Any],
    destination: str | Path,
) -> None:
    items = delivery_note["items"]
    leading = [
        ["Dokument WZ — Wydanie zewnętrzne", delivery_note.get("document_number", "")],
        ["Data wydania", delivery_note.get("date", "")],
        ["Odbiorca", delivery_note.get("customer_name", "")],
        ["NIP", delivery_note.get("customer_tax_id", "")],
        ["Zamówienie", delivery_note.get("order_number", "")],
        ["Adres dostawy", delivery_note.get("delivery_address", "")],
        ["Osoba wydająca", delivery_note.get("issuer", "")],
        ["Osoba odbierająca", delivery_note.get("receiver", "")],
        ["Wydano opakowania", delivery_note.get("issued_packages", "")],
        ["Zdano opakowania", delivery_note.get("returned_packages", "")],
        [],
    ]
    rows = [
        [
            index,
            item.get("product", ""),
            item.get("description", ""),
            item.get("description2", ""),
            item.get("quantity", ""),
            item.get("unit", ""),
            item.get("notes", ""),
        ]
        for index, item in enumerate(items, start=1)
    ]
    _write_xlsx(
        destination,
        "WZ",
        ["Lp.", "Produkt", "Opis", "Opis 2", "Ilość", "jm", "Uwagi"],
        rows,
        {0, 4},
        widths=[24, 28, 32, 32, 12, 8, 30],
        leading_rows=leading,
    )


def _render_estimate_pdf(
    estimate: dict[str, Any],
    settings: dict[str, str],
    destination: str | Path | BinaryIO,
    document_kind: str,
) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape
    except ImportError as error:
        raise ImportError("Eksport PDF wymaga biblioteki ReportLab. Zainstaluj zależności poleceniem: python -m pip install -r requirements.txt") from error

    if document_kind not in {"offer", "internal"}:
        raise ValueError("Nieznany rodzaj dokumentu PDF.")

    header = estimate["header"]
    items = estimate["items"]
    navy = colors.HexColor(settings.get("brand_primary", "#243F82"))
    red = colors.HexColor(settings.get("brand_accent", "#ED1B24"))
    ink = colors.HexColor("#18345B")
    muted = colors.HexColor("#626262")
    border = colors.HexColor("#DADADA")
    page_size = A4 if document_kind == "offer" else landscape(A4)
    output_target = str(destination) if isinstance(destination, (str, Path)) else destination
    doc = SimpleDocTemplate(
        output_target,
        pagesize=page_size,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=29 * mm,
        bottomMargin=18 * mm,
        title=f"{header['quote_number']} — BlachCalc",
        author=settings.get("company_name") or "BlachCalc",
    )

    font_regular = "Helvetica"
    font_bold = "Helvetica-Bold"
    content_width = doc.width - 12  # ramka SimpleDocTemplate ma 6 pt marginesu z każdej strony
    font_directory = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    font_candidates = (
        (font_directory / "arial.ttf", font_directory / "arialbd.ttf", font_directory / "ariali.ttf", font_directory / "arialbi.ttf"),
        (font_directory / "segoeui.ttf", font_directory / "segoeuib.ttf", font_directory / "segoeuii.ttf", font_directory / "segoeuiz.ttf"),
        (font_directory / "tahoma.ttf", font_directory / "tahomabd.ttf", font_directory / "tahoma.ttf", font_directory / "tahomabd.ttf"),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf")),
    )
    for regular_font, bold_font, italic_font, bold_italic_font in font_candidates:
        if not (regular_font.is_file() and bold_font.is_file()):
            continue
        try:
            pdfmetrics.registerFont(TTFont("BlachCalcArial", str(regular_font)))
            pdfmetrics.registerFont(TTFont("BlachCalcArialBold", str(bold_font)))
            pdfmetrics.registerFont(TTFont("BlachCalcArialItalic", str(italic_font if italic_font.is_file() else regular_font)))
            pdfmetrics.registerFont(TTFont("BlachCalcArialBoldItalic", str(bold_italic_font if bold_italic_font.is_file() else bold_font)))
            # Bez rodziny znaczniki <b>/<i> w Paragraph wracają do Helvetica, która nie ma polskich liter.
            pdfmetrics.registerFontFamily(
                "BlachCalcArial",
                normal="BlachCalcArial",
                bold="BlachCalcArialBold",
                italic="BlachCalcArialItalic",
                boldItalic="BlachCalcArialBoldItalic",
            )
            font_regular = "BlachCalcArial"
            font_bold = "BlachCalcArialBold"
            break
        except (OSError, ValueError):
            continue
    else:
        raise OSError(
            "Nie znaleziono czcionki z polskimi znakami (Arial, Segoe UI, Tahoma lub DejaVu Sans). "
            "Wydruk bez niej zawierałby puste kwadraty zamiast liter."
        )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "BlachCalcTitle", parent=styles["Title"], fontName=font_bold,
        fontSize=19, leading=23, textColor=ink, spaceAfter=3,
    )
    subtitle_style = ParagraphStyle(
        "BlachCalcSubtitle", parent=styles["Normal"], fontName=font_regular,
        fontSize=10.5, leading=15, textColor=muted,
    )
    body_style = ParagraphStyle(
        "BlachCalcBody", parent=styles["Normal"], fontName=font_regular,
        fontSize=11.5, leading=16, textColor=ink,
    )
    cell_style = ParagraphStyle(
        "BlachCalcCell", parent=body_style, fontSize=9.5, leading=12,
    )
    header_cell_style = ParagraphStyle(
        "BlachCalcHeaderCell", parent=cell_style, textColor=colors.white,
        fontName=font_bold,
    )
    cell_right_style = ParagraphStyle(
        "BlachCalcRight", parent=cell_style, alignment=TA_RIGHT,
    )
    heading_style = ParagraphStyle(
        "BlachCalcHeading", parent=body_style, fontName=font_bold,
        fontSize=12.5, leading=15, spaceBefore=7, spaceAfter=4,
    )
    header_center_style = ParagraphStyle("BlachCalcHeaderCenter", parent=header_cell_style, alignment=TA_CENTER)
    header_right_style = ParagraphStyle("BlachCalcHeaderRight", parent=header_cell_style, alignment=TA_RIGHT)

    def offer_header_row(labels: tuple[str, ...]) -> list[Any]:
        alignments = (header_center_style, header_cell_style, header_center_style, header_center_style,
                      header_right_style, header_right_style, header_right_style)
        return [Paragraph(f"<b>{escape(label)}</b>", style) for label, style in zip(labels, alignments)]
    story: list[Any] = []
    logo_path = settings.get("logo_path", "")
    company_name = settings.get("company_name", "").strip() or "Izo-Technika Iłczyk Sp. Jawna"
    company_details = [
        settings.get("company_address", "").strip(),
        f"NIP: {settings['company_tax_id'].strip()}" if settings.get("company_tax_id", "").strip() else "",
        "  |  ".join(
            value for value in (
                settings.get("company_phone", "").strip(),
                settings.get("company_email", "").strip(),
            ) if value
        ),
        settings.get("company_website", "").strip(),
    ]
    logo_file: Path | None = None
    if logo_path:
        logo_file = Path(logo_path)
        if not logo_file.is_file():
            raise ValueError(f"Nie znaleziono pliku logo: {logo_file}")
        try:
            ImageReader(str(logo_file)).getSize()
        except (OSError, ValueError) as error:
            raise ValueError(f"Nie można odczytać wybranego pliku logo: {error}") from error
    story.append(Spacer(1, 2 * mm))

    document_title = "OFERTA HANDLOWA" if document_kind == "offer" else "KARTA KALKULACJI KOSZTÓW"
    story.append(Paragraph(document_title, ParagraphStyle(
        "BlachCalcDocumentKind", parent=subtitle_style, fontName=font_bold,
        fontSize=16.5, leading=20, textColor=navy, spaceAfter=3,
    )))
    story.append(Spacer(1, 6 * mm))

    created_at = datetime.fromisoformat(header["created_at"]).astimezone()
    client_details = [
        f"<b>Zleceniodawca:</b> {escape(header['customer_name'])}",
    ]
    if str(header.get("order_number", "")).strip():
        client_details.append(f"<b>Numer zamówienia:</b> {escape(str(header['order_number']).strip())}")
    for key, label in (
        ("customer_contact_person", "Osoba kontaktowa"),
        ("customer_address", "Adres"),
        ("customer_tax_id", "NIP"),
        ("customer_email", "E-mail"),
        ("customer_phone", "Telefon"),
    ):
        value = str(header.get(key, "")).strip()
        if value:
            client_details.append(f"<b>{label}:</b> {escape(value)}")
    left_block = Paragraph("<br/>".join(client_details), body_style)
    right_block = Paragraph(
        f"<b>Numer wyceny:</b> {escape(header['quote_number'])}<br/>"
        f"<b>Data:</b> {created_at.strftime('%d.%m.%Y %H:%M')}",
        ParagraphStyle("BlachCalcMetadata", parent=body_style, alignment=TA_RIGHT),
    )
    metadata = Table([[left_block, right_block]], colWidths=[content_width * 0.61, content_width * 0.39])
    metadata.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F6F6F6")),
        ("BOX", (0, 0), (-1, -1), 0.5, border),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.extend((metadata, Spacer(1, 7 * mm)))

    vat_lines: dict[str, Decimal] = {}
    for item in items:
        vat_key = f"{Decimal(item['vat_percent']):g}%"
        vat_lines[vat_key] = vat_lines.get(vat_key, Decimal("0")) + Decimal(item["line_vat"])

    photo_width = photo_width_mm(settings) * mm

    def preview_image(item: dict[str, Any]) -> Any:
        image_data = item.get("preview_image")
        if not image_data:
            return ""
        if isinstance(image_data, memoryview):
            image_data = image_data.tobytes()
        try:
            return Image(
                BytesIO(bytes(image_data)),
                width=photo_width,
                height=photo_width * 32 / 44,
                kind="proportional",
            )
        except (OSError, TypeError, ValueError) as error:
            raise ValueError(
                f"Nie można odczytać zdjęcia pozycji „{item['element_name']}”: {error}"
            ) from error

    if document_kind == "offer":
        if header.get("pricing_mode", "itemized") == "total_only":
            headers = ("Lp.", "Zakres oferty", "Zdjęcie", "Ilość", "Wartość netto", "Wartość brutto")
            rows: list[list[Any]] = [offer_header_row(headers)]
            total_quantity = sum(int(item["quantity"]) for item in items)
            rows.append([
                "1",
                Paragraph("Kompleksowa realizacja wypalanych elementów", cell_style),
                preview_image(next((item for item in items if item.get("preview_image")), {})),
                str(total_quantity),
                Paragraph(f"{pl(Decimal(header['net_total']))} zł", cell_right_style),
                Paragraph(f"<b>{pl(Decimal(header['gross_total']))} zł</b>", cell_right_style),
            ])
            widths = [9 * mm, 40 * mm, photo_width + 4 * mm, 14 * mm, 24 * mm, 26 * mm]
            heading = "Cena całego zamówienia"
        else:
            headers = ("Lp.", "Element / zakres", "Zdjęcie", "Ilość", "Masa szt.", "Netto", "Brutto")
            rows = [offer_header_row(headers)]
            for index, item in enumerate(items, start=1):
                whole_order = item.get("calculation_mode", "per_piece") == "whole_order"
                description = escape(str(item["element_name"]))
                operation = OPERATIONS.get(item.get("operation", DEFAULT_OPERATION), "")
                if operation:
                    description += f"<br/><font color='#243F82'><b>{escape(operation)}</b></font>"
                if item.get("material"):
                    description += f"<br/><font color='#626262'>{escape(str(item['material']))}"
                    if item.get("thickness_mm"):
                        description += f" · {escape(str(item['thickness_mm']))} mm"
                    description += "</font>"
                rows.append([
                    str(index),
                    Paragraph(description, cell_style),
                    preview_image(item),
                    str(item["quantity"]),
                    "—" if whole_order else f"{Decimal(item['used_weight_kg']):.3f}".replace(".", ",") + " kg",
                    Paragraph(f"{pl(Decimal(item['line_net']))} zł", cell_right_style),
                    Paragraph(f"<b>{pl(Decimal(item['line_gross']))} zł</b>", cell_right_style),
                ])
            widths = [9 * mm, 40 * mm, photo_width + 4 * mm, 14 * mm, 19 * mm, 23 * mm, 25 * mm]
            heading = "Zestawienie pozycji oferty"
    else:
        headers = ("Element / materiał", "Ilość / zakres", "Cena zł/kg", "Masa kg/szt. lub razem", "Czas/szt. lub razem", "Stawka/godz.", "Materiał razem", "Palenie razem", "Netto razem", "Brutto razem")
        rows = [[Paragraph(f"<b>{escape(value)}</b>", header_cell_style) for value in headers]]
        for item in items:
            whole_order = item.get("calculation_mode", "per_piece") == "whole_order"
            quantity = int(item["quantity"])
            factor = 1 if whole_order else quantity
            description = escape(str(item["element_name"]))
            details = " · ".join(
                part for part in (
                    OPERATIONS.get(item.get("operation", DEFAULT_OPERATION), ""),
                    str(item.get("material", "")).strip(),
                    f"{item['thickness_mm']} mm" if item.get("thickness_mm") else "",
                ) if part
            )
            if details:
                description += f"<br/><font color='#626262'>{escape(details)}</font>"
            rows.append([
                Paragraph(description, cell_style),
                str(item["quantity"]),
                f"{Decimal(item['price_per_kg']):.2f}",
                f"{Decimal(item['used_weight_kg']):.3f}",
                _duration_label(int(item["duration_seconds"])),
                f"{Decimal(item['burning_rate_per_hour']):.2f}",
                f"{Decimal(item['material_cost']) * factor:.2f} zł",
                f"{Decimal(item['burning_cost']) * factor:.2f} zł",
                f"{Decimal(item['line_net']):.2f} zł",
                f"{Decimal(item['line_gross']):.2f} zł",
            ])
        widths = [43 * mm, 11 * mm, 17 * mm, 18 * mm, 20 * mm, 19 * mm, 23 * mm, 22 * mm, 25 * mm, 25 * mm]
        heading = "Szczegółowa kalkulacja kosztów"

    story.append(Paragraph(heading, heading_style))
    if document_kind == "offer":
        if not any(item.get("preview_image") for item in items):
            rows = [row[:2] + row[3:] for row in rows]
            widths = widths[:2] + widths[3:]
        widths[1] = content_width - (sum(widths) - widths[1])
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), font_bold),
        ("FONTNAME", (0, 1), (-1, -1), font_regular),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.35, border),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), (colors.white, colors.HexColor("#FAFAFA"))),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT" if document_kind == "internal" else "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.append(table)

    net_total = Decimal(header["net_total"])
    gross_total = Decimal(header["gross_total"])
    money = pl if document_kind == "offer" else (lambda value: f"{value:.2f}")
    totals = [
        ["Razem netto", f"{money(net_total)} zł"],
        *[[f"VAT {rate}", f"{money(amount)} zł"] for rate, amount in vat_lines.items()],
        ["Razem brutto", f"{money(gross_total)} zł"],
    ]
    totals_table = Table(
        [[Paragraph(f"<b>{escape(label)}</b>", body_style), Paragraph(f"<b>{escape(amount)}</b>", cell_right_style)] for label, amount in totals],
        colWidths=[38 * mm, 35 * mm],
        hAlign="RIGHT",
    )
    totals_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 0.5, border),
        ("LINEABOVE", (0, -1), (-1, -1), 1.1, red),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F0F3FA")),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(KeepTogether([Spacer(1, 5 * mm), totals_table]))

    offer_conditions = []
    validity_days = settings.get("offer_validity_days", "").strip()
    if validity_days:
        valid_until = ""
        if validity_days.isdigit():
            valid_until = f" (do {(created_at + timedelta(days=int(validity_days))).strftime('%d.%m.%Y')})"
        offer_conditions.append(f"<b>Ważność oferty:</b> {escape(validity_days)} dni od daty wystawienia{valid_until}")
    for key, label in (
        ("payment_terms", "Warunki płatności"),
        ("delivery_terms", "Termin realizacji"),
    ):
        value = settings.get(key, "").strip()
        if value:
            offer_conditions.append(f"<b>{label}:</b> {escape(value)}")
    if document_kind == "offer" and offer_conditions:
        story.extend((
            Spacer(1, 5 * mm),
            Paragraph("Warunki oferty", heading_style),
            Paragraph("<br/>".join(offer_conditions), body_style),
        ))

    notes = str(header.get("notes", "")).strip()
    if notes:
        story.extend((Spacer(1, 5 * mm), Paragraph("Uwagi", heading_style), Paragraph(escape(notes).replace("\n", "<br/>"), body_style)))

    if document_kind == "offer":
        terms_style = ParagraphStyle(
            "BlachCalcOfferTerms", parent=body_style, fontSize=9,
            leading=11.6, leftIndent=9, firstLineIndent=-7, spaceAfter=2,
        )
        page2 = parse_offer_page(settings.get(OFFER_PAGE_SETTING, ""))
        section_style = ParagraphStyle(
            "BlachCalcSection", parent=heading_style, textColor=ink,
            fontSize=13, leading=16, spaceBefore=0, spaceAfter=0,
        )

        def section_heading(text: str) -> Any:
            block = Table([[Paragraph(escape(text), section_style)]], colWidths=[content_width], hAlign="LEFT")
            block.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0F3FA")),
                ("LINEBEFORE", (0, 0), (0, 0), 3, red),
                ("LINEBELOW", (0, 0), (-1, 0), 0.6, navy),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            block.keepWithNext = True
            return KeepTogether([Spacer(1, 3 * mm), block, Spacer(1, 2 * mm)])

        def bullet_list(key: str) -> list[Any]:
            return [Paragraph(f"- {escape(term)}", terms_style) for term in page2[key]["items"]]

        def titled(key: str, content: list[Any]) -> list[Any]:
            title = page2[key]["title"]
            return [section_heading(title), *content] if title else content

        story.extend((PageBreak(), *titled("terms", bullet_list("terms"))))
        pallet_rows = [[line] for line in page2["pallets"]["items"]]
        pallet_table = Table(pallet_rows or [[""]], colWidths=[76 * mm], hAlign="LEFT")
        pallet_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAFAFA")),
            ("GRID", (0, 0), (-1, -1), 0.35, border),
            ("FONTNAME", (0, 0), (-1, -1), font_regular),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        contact_label_style = ParagraphStyle(
            "BlachCalcContactLabel", parent=cell_style, fontSize=8, leading=10, textColor=muted,
        )
        contact_value_style = ParagraphStyle(
            "BlachCalcContactValue", parent=cell_style, fontSize=11, leading=14, textColor=ink, fontName=font_bold,
        )
        contacts = [
            [Paragraph("OSOBA KONTAKTOWA", contact_label_style), Paragraph("TELEFON", contact_label_style), Paragraph("E-MAIL", contact_label_style)],
        ]
        contacts.extend([
            [
                Paragraph(escape(contact[field]), contact_value_style)
                for field in CONTACT_FIELDS
            ]
            for contact in page2["contacts"]
        ])
        contacts_table = Table(
            contacts,
            colWidths=[50 * mm, 40 * mm, 70 * mm],
            repeatRows=1,
            hAlign="LEFT",
        )
        contacts_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0F3FA")),
            ("LINEBEFORE", (0, 0), (0, -1), 3, red),
            ("LINEAFTER", (0, 0), (1, -1), 0.35, border),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, 0), 6),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 0),
            ("TOPPADDING", (0, 1), (-1, 1), 1),
            ("BOTTOMPADDING", (0, 1), (-1, 1), 7),
        ]))
        story.extend(titled("claims", bullet_list("claims")))
        if page2["pallets"]["items"]:
            story.extend(titled("pallets", [pallet_table]))
        for section in page2["custom_sections"]:
            if section["title"]:
                story.append(section_heading(section["title"]))
            if section["text"]:
                story.append(
                    Paragraph(
                        escape(section["text"]).replace("\n", "<br/>"),
                        body_style,
                    )
                )
        if page2["contacts"]:
            story.extend(titled("contact", [contacts_table]))

    footer_text = (
        "Dokument wewnętrzny — kalkulacja kosztów produkcyjnych."
        if document_kind == "internal"
        else "Oferta handlowa — dokument nie jest fakturą VAT."
    )

    def page_header(canvas: Any, page_number: int) -> None:
        canvas.saveState()
        logo_x = 15 * mm + 6
        logo_y = page_size[1] - 22 * mm
        if logo_file is not None:
            canvas.drawImage(
                str(logo_file),
                logo_x,
                logo_y,
                width=58 * mm,
                height=16 * mm,
                preserveAspectRatio=True,
                anchor="c",
                mask="auto",
            )
        details = [f"<b>{escape(company_name)}</b>"]
        details.extend(escape(line) for line in company_details if line)
        details_block = Paragraph(
            "<br/>".join(details),
            ParagraphStyle(
                "BlachCalcRepeatedHeader",
                parent=subtitle_style,
                alignment=TA_RIGHT,
                fontSize=7.5,
                leading=9,
                textColor=ink,
            ),
        )
        details_width = content_width - 62 * mm
        _, details_height = details_block.wrap(details_width, 20 * mm)
        details_block.drawOn(
            canvas,
            page_size[0] - 15 * mm - 6 - details_width,
            page_size[1] - 6 * mm - details_height,
        )
        canvas.setStrokeColor(red)
        canvas.setLineWidth(1.2)
        canvas.line(
            15 * mm + 6,
            page_size[1] - 25 * mm,
            page_size[0] - 15 * mm - 6,
            page_size[1] - 25 * mm,
        )
        canvas.restoreState()

    def footer(canvas: Any, page_count: int) -> None:
        canvas.saveState()
        canvas.setStrokeColor(border)
        canvas.line(15 * mm + 6, 13 * mm, page_size[0] - 15 * mm - 6, 13 * mm)
        company_footer = "  |  ".join(
            value for value in (
                company_name,
                settings.get("company_tax_id", "").strip(),
                settings.get("company_website", "").strip(),
            ) if value
        )
        footer_block = Paragraph(
            escape(f"{company_footer}  |  {footer_text}"),
            ParagraphStyle(
                "BlachCalcRepeatedFooter",
                parent=subtitle_style,
                fontSize=6.5,
                leading=7.5,
                textColor=muted,
            ),
        )
        footer_width = content_width - 27 * mm
        footer_block.wrap(footer_width, 9 * mm)
        footer_block.drawOn(canvas, 15 * mm + 6, 3 * mm)
        canvas.setFont(font_regular, 6.5)
        canvas.setFillColor(muted)
        canvas.drawRightString(page_size[0] - 15 * mm - 6, 8 * mm, f"Strona {canvas.getPageNumber()} z {page_count}")
        canvas.restoreState()

    class NumberedCanvas(Canvas):
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            super().__init__(*args, **kwargs)
            self._saved_pages: list[dict[str, Any]] = []

        def showPage(self) -> None:
            self._saved_pages.append(dict(self.__dict__))
            self._startPage()

        def save(self) -> None:
            page_count = len(self._saved_pages)
            for page_number, state in enumerate(self._saved_pages, start=1):
                self.__dict__.update(state)
                page_header(self, page_number)
                footer(self, page_count)
                super().showPage()
            super().save()

    doc.build(story, canvasmaker=NumberedCanvas)


def export_estimate_pdf(
    estimate: dict[str, Any],
    settings: dict[str, str],
    destination: str | Path | BinaryIO,
    document_kind: str,
) -> None:
    if isinstance(destination, (str, Path)):
        with _atomic_destination(destination) as temporary_path:
            _render_estimate_pdf(estimate, settings, temporary_path, document_kind)
        return

    _render_estimate_pdf(estimate, settings, destination, document_kind)


def export_delivery_note_pdf(
    delivery_note: dict[str, Any],
    settings: dict[str, str],
    destination: str | Path | BinaryIO,
) -> None:
    if not delivery_note["items"]:
        raise ValueError("Dodaj co najmniej jedną pozycję do wydania zewnętrznego.")
    for field, label in (
        ("document_number", "Numer WZ"),
        ("order_number", "Numer zamówienia"),
        ("customer_name", "Odbiorca"),
        ("delivery_address", "Adres dostawy"),
        ("date", "Data wydania"),
    ):
        if not str(delivery_note.get(field) or "").strip():
            raise ValueError(f"Uzupełnij pole: {label}.")

    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape
    except ImportError as error:
        raise ImportError("Eksport WZ wymaga biblioteki ReportLab.") from error

    page_size = landscape(A4)
    page_width, page_height = page_size
    left_margin = 15 * mm
    right_margin = 15 * mm
    content_width = page_width - left_margin - right_margin
    company_name = settings.get("company_name", "").strip() or "Izo-Technika Iłczyk Sp. Jawna"
    company_address = settings.get("company_address", "").strip()
    company_tax_id = settings.get("company_tax_id", "").strip()
    company_regon = settings.get("company_regon", "").strip()
    company_phone = settings.get("company_phone", "").strip()
    company_email = settings.get("company_email", "").strip()
    logo_path = settings.get("logo_path", "").strip()
    logo_file = Path(logo_path) if logo_path else None
    if logo_file is not None and not logo_file.is_file():
        raise ValueError(f"Nie znaleziono pliku logo: {logo_file}")

    font_regular = "Helvetica"
    font_bold = "Helvetica-Bold"
    font_directory = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    font_candidates = (
        (font_directory / "arial.ttf", font_directory / "arialbd.ttf"),
        (font_directory / "segoeui.ttf", font_directory / "segoeuib.ttf"),
        (font_directory / "tahoma.ttf", font_directory / "tahomabd.ttf"),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
    )
    for regular_font, bold_font in font_candidates:
        if not (regular_font.is_file() and bold_font.is_file()):
            continue
        try:
            pdfmetrics.registerFont(TTFont("BlachCalcWZ", str(regular_font)))
            pdfmetrics.registerFont(TTFont("BlachCalcWZBold", str(bold_font)))
            pdfmetrics.registerFontFamily(
                "BlachCalcWZ",
                normal="BlachCalcWZ",
                bold="BlachCalcWZBold",
                italic="BlachCalcWZ",
                boldItalic="BlachCalcWZBold",
            )
            font_regular = "BlachCalcWZ"
            font_bold = "BlachCalcWZBold"
            break
        except (OSError, ValueError):
            continue
    else:
        raise OSError("Nie znaleziono czcionki z polskimi znakami do wygenerowania WZ.")

    navy = colors.HexColor(settings.get("brand_primary", "#243F82"))
    accent = colors.HexColor(settings.get("brand_accent", "#ED1B24"))
    heading_fill = colors.HexColor("#DCE8F7")
    border = colors.HexColor("#395B91")
    text_style = ParagraphStyle(
        "BlachCalcWZText",
        fontName=font_regular,
        fontSize=9.5,
        leading=12,
        textColor=navy,
        alignment=TA_LEFT,
    )
    cell_style = ParagraphStyle(
        "BlachCalcWZCell",
        parent=text_style,
        fontSize=9.5,
        leading=12,
        alignment=TA_CENTER,
    )
    header_style = ParagraphStyle(
        "BlachCalcWZHeader",
        parent=cell_style,
        fontName=font_bold,
        fontSize=9.5,
        leading=12,
    )
    info_style = ParagraphStyle(
        "BlachCalcWZInfo",
        parent=text_style,
        fontSize=11.5,
        leading=16,
    )
    info_right_style = ParagraphStyle(
        "BlachCalcWZInfoRight",
        parent=info_style,
        alignment=TA_RIGHT,
    )
    signature_style = ParagraphStyle(
        "BlachCalcWZSignature",
        parent=text_style,
        fontSize=9.5,
        leading=12,
    )

    def paragraph(
        value: Any,
        style: ParagraphStyle = cell_style,
        markup: bool = False,
    ) -> Paragraph:
        safe_text = (
            str(value)
            if markup
            else escape(str(value or "")).replace("\n", "<br/>")
        )
        return Paragraph(safe_text, style)

    def build_document(output: str | Path | BinaryIO) -> None:
        document = SimpleDocTemplate(
            output,
            pagesize=page_size,
            leftMargin=left_margin,
            rightMargin=right_margin,
            topMargin=30 * mm,
            bottomMargin=18 * mm,
            title=f"{delivery_note['document_number']} — Wydanie zewnętrzne",
            author=company_name,
        )
        story: list[Any] = []
        note_items = delivery_note["items"]
        rows_per_page = 14
        column_widths = [
            8 * mm,
            33 * mm,
            44 * mm,
            44 * mm,
            18 * mm,
            14 * mm,
            content_width - 161 * mm,
        ]
        row_padding = max(1.0, (wz_row_height_mm(settings) * mm - 11.5) / 2)
        table_header = [
            paragraph(label, header_style)
            for label in ("Lp.", "Produkt", "Opis", "Opis 2", "Ilość", "jm", "Uwagi")
        ]
        for start in range(0, len(note_items), rows_per_page):
            page_items = note_items[start:start + rows_per_page]
            rows: list[list[Any]] = [table_header]
            for index, item in enumerate(page_items, start=start + 1):
                product = str(item.get("product", "")).strip()
                quantity_text = str(item.get("quantity", "")).strip().replace(",", ".")
                unit = str(item.get("unit", "")).strip()
                if not product:
                    raise ValueError(f"Pozycja {index}: uzupełnij nazwę produktu.")
                if not unit:
                    raise ValueError(f"Pozycja {index}: uzupełnij jednostkę miary.")
                try:
                    quantity = Decimal(quantity_text)
                except (InvalidOperation, ValueError) as error:
                    raise ValueError(f"Pozycja {index}: ilość musi być prawidłową liczbą.") from error
                if not quantity.is_finite() or quantity <= 0:
                    raise ValueError(f"Pozycja {index}: ilość musi być większa od zera.")
                if quantity > Decimal("1000000000"):
                    raise ValueError(f"Pozycja {index}: ilość jest zbyt duża.")
                rows.append([
                    paragraph(index),
                    paragraph(product),
                    paragraph(item.get("description", "")),
                    paragraph(item.get("description2", "")),
                    paragraph(item["quantity"]),
                    paragraph(unit),
                    paragraph(item.get("notes", "")),
                ])
            for _ in range(rows_per_page - len(page_items)):
                rows.append([paragraph(" ") for _ in table_header])
            table = Table(
                rows,
                colWidths=column_widths,
                repeatRows=1,
                hAlign="LEFT",
                splitByRow=1,
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), heading_fill),
                ("TEXTCOLOR", (0, 0), (-1, -1), navy),
                ("GRID", (0, 0), (-1, -1), 0.55, border),
                ("FONTNAME", (0, 0), (-1, 0), font_bold),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), row_padding),
                ("BOTTOMPADDING", (0, 0), (-1, -1), row_padding),
                ("LINEBELOW", (0, 0), (-1, 0), 1, navy),
            ]))
            story.append(table)
            if start + rows_per_page < len(note_items):
                story.append(PageBreak())

        metadata = Table(
            [[
                paragraph(
                    f"<b>Numer dokumentu:</b> {escape(delivery_note['document_number'])}<br/>"
                    f"<b>Odbiorca:</b> {escape(delivery_note['customer_name'])}<br/>"
                    f"<b>NIP:</b> {escape(delivery_note.get('customer_tax_id', ''))}",
                    info_style,
                    markup=True,
                ),
                paragraph(
                    f"<b>Zamówienie:</b> {escape(delivery_note['order_number'])}<br/>"
                    f"<b>Adres dostawy:</b> {escape(delivery_note['delivery_address'])}",
                    info_right_style,
                    markup=True,
                ),
            ]],
            colWidths=[content_width * 0.49, content_width * 0.51],
        )
        metadata.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.insert(0, metadata)
        story.insert(0, paragraph(
            f"<b>Data wydania:</b> {escape(delivery_note['date'])}",
            info_right_style,
            markup=True,
        ))
        story.insert(0, Spacer(1, 3 * mm))

        signature_line_style = ParagraphStyle(
            "BlachCalcWZSignatureLine",
            parent=signature_style,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#5B6B84"),
            fontSize=8,
            leading=10,
        )
        story.extend((
            Spacer(1, 5 * mm),
            Table(
                [
                    [
                        paragraph(
                            f"<b>Osoba wydająca:</b> {escape(delivery_note.get('issuer', ''))}<br/>"
                            "Potwierdzam wydanie w/w towarów bez zastrzeżeń.",
                            signature_style,
                            markup=True,
                        ),
                        paragraph(
                            f"<b>Osoba odbierająca:</b> {escape(delivery_note.get('receiver', ''))}<br/>"
                            "Potwierdzam odbiór w/w towarów bez zastrzeżeń.",
                            signature_style,
                            markup=True,
                        ),
                    ],
                    [
                        paragraph("podpis i data", signature_line_style),
                        paragraph("podpis i data", signature_line_style),
                    ],
                ],
                colWidths=[content_width / 2, content_width / 2],
                rowHeights=[None, 16 * mm],
                style=TableStyle([
                    ("BOX", (0, 0), (-1, -1), 0.55, border),
                    ("LINEAFTER", (0, 0), (0, -1), 0.55, border),
                    ("VALIGN", (0, 0), (-1, 0), "TOP"),
                    ("VALIGN", (0, 1), (-1, 1), "BOTTOM"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]),
            ),
            Spacer(1, 5 * mm),
        ))

        packaging = Table(
            [[
                paragraph(
                    "<b>Wydano opakowania:</b><br/>"
                    + escape(str(delivery_note.get("issued_packages", "")).strip()).replace("\n", "<br/>"),
                    signature_style,
                    markup=True,
                ),
                paragraph(
                    "<b>Zdano opakowania:</b><br/>"
                    + escape(str(delivery_note.get("returned_packages", "")).strip()).replace("\n", "<br/>"),
                    signature_style,
                    markup=True,
                ),
            ]],
            colWidths=[content_width / 2, content_width / 2],
        )
        packaging.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.55, border),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(packaging)

        def page_chrome(canvas: Any, page_number: int, page_count: int) -> None:
            canvas.saveState()
            company_lines = [company_name]
            if company_address:
                company_lines.append(f"Adres: {company_address}")
            identifiers = "  |  ".join(
                f"{label}: {value}"
                for label, value in (("NIP", company_tax_id), ("REGON", company_regon))
                if value
            )
            contacts = "  |  ".join(
                f"{label}: {value}"
                for label, value in (("Tel", company_phone), ("e-mail", company_email))
                if value
            )
            company_lines.extend(value for value in (identifiers, contacts) if value)
            company_block = Paragraph(
                "<br/>".join(escape(line) for line in company_lines),
                ParagraphStyle(
                    "BlachCalcWZCompany",
                    parent=text_style,
                    fontName=font_bold,
                    fontSize=8,
                    leading=10,
                ),
            )
            company_block.wrap(145 * mm, 22 * mm)
            company_block.drawOn(canvas, left_margin, page_height - 6 * mm - company_block.height)
            canvas.setFont(font_regular, 8)
            canvas.setFillColor(navy)
            canvas.drawRightString(page_width - right_margin, page_height - 7 * mm, delivery_note["date"])
            if logo_file is not None:
                canvas.drawImage(
                    str(logo_file),
                    page_width - right_margin - 66 * mm,
                    page_height - 26 * mm,
                    width=66 * mm,
                    height=19 * mm,
                    preserveAspectRatio=True,
                    anchor="c",
                    mask="auto",
                )
            canvas.setFillColor(colors.HexColor("#5485BE"))
            canvas.rect(left_margin, page_height - 29 * mm, content_width, 16, stroke=0, fill=1)
            canvas.setFillColor(colors.white)
            canvas.setFont(font_bold, 10)
            canvas.drawCentredString(
                page_width / 2,
                page_height - 29 * mm + 4.5,
                "Dokument WZ - Wydanie Zewnętrzne",
            )
            canvas.setStrokeColor(colors.HexColor("#B9C5D5"))
            canvas.line(left_margin, 13 * mm, page_width - right_margin, 13 * mm)
            canvas.setFillColor(navy)
            canvas.setFont(font_regular, 7)
            canvas.drawString(left_margin, 8 * mm, f"{company_name}  |  {delivery_note['document_number']}")
            canvas.drawRightString(
                page_width - right_margin,
                8 * mm,
                f"Strona {page_number} z {page_count}",
            )
            canvas.restoreState()

        class NumberedCanvas(Canvas):
            def __init__(self, *args: Any, **kwargs: Any) -> None:
                super().__init__(*args, **kwargs)
                self._pages: list[dict[str, Any]] = []

            def showPage(self) -> None:
                self._pages.append(dict(self.__dict__))
                self._startPage()

            def save(self) -> None:
                page_count = len(self._pages)
                for page_number, state in enumerate(self._pages, start=1):
                    self.__dict__.update(state)
                    page_chrome(self, page_number, page_count)
                    super().showPage()
                super().save()

        document.build(story, canvasmaker=NumberedCanvas)

    if isinstance(destination, (str, Path)):
        with _atomic_destination(destination) as temporary_path:
            build_document(str(temporary_path))
    else:
        build_document(destination)


def _duration_label(seconds: int) -> str:
    hours, remainder = divmod(seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours:02}:{minutes:02}:{seconds:02}"
